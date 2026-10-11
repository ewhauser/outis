"""Native fixtures contain synthetic identifiers; no models or external files."""

import base64
import io
import json
import posixpath
import re
import zipfile
from xml.etree import ElementTree as ET

import pytest

from outis.formats.files import MIME_TYPES, A, C, P, R, S, W, process_file


class StubService:
    version = "test-model"

    def process_segments(self, requests, **kwargs):
        results = []
        for request in requests:
            text = request["text"]
            matches = list(re.finditer(r"Jane Doe|Zoë Núñez|2024-10-03", text))
            entities = [
                {
                    "start": m.start(),
                    "end": m.end(),
                    "label": "NAME",
                    "replacement": "[NAME]",
                }
                for m in matches
            ]
            results.append(
                {
                    "text": re.sub(r"Jane Doe|Zoë Núñez|2024-10-03", "[NAME]", text),
                    "entities": entities,
                    "model_version": self.version,
                }
            )
        return results


def encode(raw):
    return base64.b64encode(raw).decode("ascii")


def invoke(raw, kind, **kwargs):
    return process_file(
        StubService(), data=encode(raw), content_type=MIME_TYPES[kind], **kwargs
    )


def make_office(kind, parts=None):
    mains = {
        "docx": "word/document.xml",
        "xlsx": "xl/workbook.xml",
        "pptx": "ppt/presentation.xml",
    }
    default = {
        "docx": {
            "word/document.xml": f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Jane </w:t></w:r><w:r><w:t>Doe</w:t></w:r></w:p></w:body></w:document>'
        },
        "xlsx": {
            "xl/workbook.xml": f'<workbook xmlns="{S}"><sheets><sheet name="Patients" sheetId="1"/></sheets></workbook>',
            "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{S}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><r><t>Jane </t></r><r><t>Doe</t></r></is></c></row></sheetData></worksheet>',
        },
        "pptx": {
            "ppt/presentation.xml": f'<p:presentation xmlns:p="{P}"/>',
            "ppt/slides/slide1.xml": f'<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Jane </a:t></a:r><a:r><a:t>Doe</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>',
        },
    }[kind]
    members = {
        "[Content_Types].xml": f'<Types xmlns="{C}"><Override PartName="/docProps/core.xml" ContentType="core"/></Types>',
        "_rels/.rels": f'<Relationships xmlns="{R}"><Relationship Id="r1" Type="officeDocument" Target="{mains[kind]}"/><Relationship Id="r2" Type="metadata" Target="docProps/core.xml"/></Relationships>',
        "docProps/core.xml": "<properties><creator>Jane Doe</creator></properties>",
        **default,
        **(parts or {}),
    }
    return zip_bytes(members)


def zip_bytes(members):
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return raw.getvalue()


def unpack(result):
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(result["data"]))) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_split_native_runs_and_metadata(kind):
    result = invoke(make_office(kind), kind)
    members = unpack(result)
    assert "docProps/core.xml" not in members
    assert result["entity_count"] == 1
    assert all(
        "Jane" not in raw.decode() and "Doe" not in raw.decode()
        for raw in members.values()
    )
    assert any("[NAME]" in raw.decode() for raw in members.values())
    assert isinstance(result["entities"][0]["location"]["part"], int)
    for raw in members.values():
        ET.fromstring(raw)


def test_docx_header_footer_and_footnotes_processed():
    parts = {
        f"word/{part}.xml": f'<w:root xmlns:w="{W}"><w:p><w:r><w:t>Jane Doe</w:t></w:r></w:p></w:root>'
        for part in ["header1", "footer1", "footnotes"]
    }
    result = invoke(make_office("docx", parts), "docx")
    assert result["entity_count"] == 4
    assert all(b"Jane" not in raw for raw in unpack(result).values())


def shape(element):
    return [(child.tag.rsplit("}", 1)[-1], child.text or "") for child in element]


def test_docx_breaks_do_not_join_words():
    part = f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Jane</w:t><w:tab/><w:t>Doe</w:t></w:r><w:r><w:br w:type="textWrapping"/><w:t>Seen by Jane Doe</w:t><w:cr/><w:t>daily</w:t></w:r></w:p></w:body></w:document>'
    seen = []
    service = StubService()
    detect = service.process_segments
    service.process_segments = lambda requests, **kw: (
        seen.extend(r["text"] for r in requests) or detect(requests, **kw)
    )
    result = process_file(
        service,
        data=encode(make_office("docx", {"word/document.xml": part})),
        content_type=MIME_TYPES["docx"],
    )
    assert "Jane\tDoe\nSeen by Jane Doe\ndaily" in seen
    root = ET.fromstring(unpack(result)["word/document.xml"])
    first, second = root.iter(f"{{{W}}}r")
    # Word does not render tabs or line feeds inside w:t; they stay elements.
    assert shape(first) == [
        ("t", "Jane"),
        ("tab", ""),
        ("t", "Doe"),
        ("br", ""),
        ("t", "Seen by [NAME]"),
        ("cr", ""),
        ("t", "daily"),
        ("t", ""),
    ]
    assert first[3].get(f"{{{W}}}type") == "textWrapping"
    assert shape(second) == [("t", ""), ("t", "")]


def test_docx_inline_characters_stay_in_place():
    part = f'<w:document xmlns:w="{W}"><w:body><w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr><w:r><w:t>Dose: 1</w:t><w:noBreakHyphen/><w:t>2 mg Jane</w:t><w:softHyphen/></w:r><w:r><w:t> Doe</w:t></w:r></w:p></w:body></w:document>'
    result = invoke(make_office("docx", {"word/document.xml": part}), "docx")
    root = ET.fromstring(unpack(result)["word/document.xml"])
    assert "".join(root.itertext()) == "Dose: 1\u20112 mg [NAME]"
    assert root.find(f".//{{{W}}}noBreakHyphen") is None
    assert root.find(f".//{{{W}}}softHyphen") is None
    # Paragraph tab stops are formatting, not run content.
    assert root.find(f".//{{{W}}}tabs/{{{W}}}tab") is not None


def test_docx_dynamic_run_content_rejected():
    part = f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Seen </w:t><w:dayLong/></w:r></w:p></w:body></w:document>'
    with pytest.raises(ValueError, match="dynamic"):
        invoke(make_office("docx", {"word/document.xml": part}), "docx")


def test_pptx_soft_breaks_do_not_join_words():
    slide = f'<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:rPr b="1"/><a:t>Take 1</a:t></a:r><a:br><a:rPr b="1"/></a:br><a:r><a:t>tablet for Jane </a:t></a:r><a:r><a:t>Doe</a:t></a:r><a:endParaRPr/></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
    seen = []
    service = StubService()
    detect = service.process_segments
    service.process_segments = lambda requests, **kw: (
        seen.extend(r["text"] for r in requests) or detect(requests, **kw)
    )
    result = process_file(
        service,
        data=encode(make_office("pptx", {"ppt/slides/slide1.xml": slide})),
        content_type=MIME_TYPES["pptx"],
    )
    assert "Take 1\ntablet for Jane Doe" in seen
    root = ET.fromstring(unpack(result)["ppt/slides/slide1.xml"])
    paragraph = next(root.iter(f"{{{A}}}p"))
    # PowerPoint does not render a line feed inside a:t; it stays an a:br.
    assert shape(paragraph) == [
        ("r", ""),
        ("br", ""),
        ("r", ""),
        ("r", ""),
        ("r", ""),
        ("endParaRPr", ""),
    ]
    assert [shape(run) for run in paragraph.iter(f"{{{A}}}r")] == [
        [("rPr", ""), ("t", "Take 1")],
        [("rPr", ""), ("t", "tablet for [NAME]")],
        [("t", "")],
        [("t", "")],
    ]
    assert paragraph[1].find(f"{{{A}}}rPr") is not None


@pytest.mark.parametrize(
    "code", ["m", "mm", "mmm", "mmmm", "mmmmm", "[$-409]mmmm", "bbbb"]
)
def test_xlsx_month_only_formats_are_dates(code):
    parts = {
        "xl/styles.xml": f'<styleSheet xmlns="{S}"><numFmts><numFmt numFmtId="166" formatCode="{code}"/></numFmts><cellXfs><xf numFmtId="0"/><xf numFmtId="166"/></cellXfs></styleSheet>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{S}"><sheetData><row><c r="A1" s="1"><v>45000</v></c></row></sheetData></worksheet>',
    }
    members = unpack(invoke(make_office("xlsx", parts), "xlsx"))
    sheet = members["xl/worksheets/sheet1.xml"].decode()
    assert "45000" not in sheet
    assert "2023-03-15" in sheet


@pytest.mark.parametrize("code", ["General", "0.00", "0.00E+00", "#,##0;[Red]-#,##0"])
def test_xlsx_numeric_formats_are_not_dates(code):
    parts = {
        "xl/styles.xml": f'<styleSheet xmlns="{S}"><numFmts><numFmt numFmtId="166" formatCode="{code}"/></numFmts><cellXfs><xf numFmtId="0"/><xf numFmtId="166"/></cellXfs></styleSheet>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{S}"><sheetData><row><c r="A1" s="1"><v>45000</v></c></row></sheetData></worksheet>',
    }
    members = unpack(invoke(make_office("xlsx", parts), "xlsx"))
    assert "45000" in members["xl/worksheets/sheet1.xml"].decode()


@pytest.mark.parametrize(
    "part,body",
    [
        ("word/comments.xml", "<comments/>"),
        ("word/media/image1.png", b"image"),
        ("word/vbaProject.bin", b"binary"),
    ],
)
def test_unsupported_parts_fail_closed(part, body):
    with pytest.raises(ValueError, match="unsupported parts"):
        invoke(make_office("docx", {part: body}), "docx")


@pytest.mark.parametrize(
    "body",
    [
        f'<w:fldSimple xmlns:w="{W}" w:instr="Jane Doe"/>',
        f'<w:ins xmlns:w="{W}"><w:t>Jane Doe</w:t></w:ins>',
    ],
)
def test_fields_and_changes_rejected(body):
    with pytest.raises(ValueError, match="unsupported fields"):
        invoke(make_office("docx", {"word/settings.xml": body}), "docx")


def test_external_relationships_rejected():
    part = f'<Relationships xmlns="{R}"><Relationship Id="r1" Target="https://example.invalid/patient" TargetMode="External"/></Relationships>'
    with pytest.raises(ValueError, match="External"):
        invoke(make_office("docx", {"word/_rels/document.xml.rels": part}), "docx")


def test_xlsx_formulas_and_hidden_rows_rejected():
    for content in [
        "<row><c><f>Jane Doe</f><v>1</v></c></row>",
        '<row hidden="1"><c/></row>',
    ]:
        part = f'<worksheet xmlns="{S}"><sheetData>{content}</sheetData></worksheet>'
        with pytest.raises(ValueError):
            invoke(make_office("xlsx", {"xl/worksheets/sheet1.xml": part}), "xlsx")


def test_xlsx_shared_strings_dates_and_names():
    parts = {
        "xl/sharedStrings.xml": f'<sst xmlns="{S}"><si><r><t>Jane </t></r><r><t>Doe</t></r></si></sst>',
        "xl/styles.xml": f'<styleSheet xmlns="{S}"><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{S}"><sheetData><row><c r="A1" t="s"><v>0</v></c><c r="A2" s="1"><v>45568</v></c></row></sheetData></worksheet>',
        "xl/workbook.xml": f'<workbook xmlns="{S}"><sheets><sheet name="Jane Doe" sheetId="1"/></sheets></workbook>',
    }
    result = invoke(make_office("xlsx", parts), "xlsx")
    assert result["entity_count"] == 3
    members = unpack(result)
    assert all(b"Jane" not in raw for raw in members.values())
    assert b"45568" not in members["xl/worksheets/sheet1.xml"]


@pytest.mark.parametrize(
    "path", ["../secret.xml", "/secret.xml", "a/../secret.xml", "a\\secret.xml"]
)
def test_zip_traversal_rejected(path):
    with pytest.raises(ValueError, match="unsafe"):
        invoke(zip_bytes({path: "secret"}), "docx")


def test_zip_bomb_and_duplicate_rejected():
    with pytest.raises(ValueError, match="compression ratio"):
        invoke(zip_bytes({"word/document.xml": b"a" * 2_000_000}), "docx")
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as archive:
        archive.writestr("a", "1")
        with pytest.warns(UserWarning):
            archive.writestr("a", "2")
    with pytest.raises(ValueError, match="duplicate"):
        invoke(raw.getvalue(), "docx")


def test_unsafe_office_xml_rejected():
    part = (
        '<!DOCTYPE data [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><data>&xxe;</data>'
    )
    with pytest.raises(ValueError, match="unsafe XML"):
        invoke(make_office("docx", {"word/settings.xml": part}), "docx")


def test_utf8_text_and_json_csv_xml_files():
    fixtures = {
        "txt": "Jane Doe",
        "json": '{"name": "Jane Doe"}',
        "csv": "name\nJane Doe\n",
        "xml": "<record>Jane Doe</record>",
    }
    for kind, text in fixtures.items():
        result = invoke(text.encode(), kind)
        output = base64.b64decode(result["data"]).decode()
        assert "Jane" not in output and "[NAME]" in output
        assert result["entities"][0]["location"]
        assert result["entity_count"] == 1
    assert json.loads(
        base64.b64decode(invoke(fixtures["json"].encode(), "json")["data"])
    ) == {"name": "[NAME]"}


@pytest.mark.parametrize(
    "raw", [b'{"x": "Jane Doe", "x": "duplicate"}', b'{"x": NaN}', b"not-json"]
)
def test_invalid_json_rejected(raw):
    with pytest.raises(ValueError, match="valid JSON"):
        invoke(raw, "json")


def test_invalid_type_encoding_and_paths():
    with pytest.raises(ValueError, match="Unsupported content"):
        process_file(StubService(), data="aGVsbG8=", content_type="image/png")
    with pytest.raises(ValueError, match="base64"):
        process_file(StubService(), data="not-base64!", content_type="txt")
    with pytest.raises(ValueError, match="UTF-8"):
        invoke(b"\xff\xfe\xfa", "txt")
    with pytest.raises(ValueError, match="Field selection"):
        invoke(b"Jane Doe", "txt", paths=["/name"])


def make_pdf(text="Jane Doe", blank=False):
    from reportlab.pdfgen.canvas import Canvas

    buf = io.BytesIO()
    canvas = Canvas(buf)
    canvas.setAuthor("Jane Doe")
    if not blank:
        canvas.drawString(20, 750, text)
    canvas.showPage()
    canvas.save()
    return buf.getvalue()


def test_digital_pdf_extract_redact_rebuild():
    from pypdf import PdfReader

    result = invoke(make_pdf("Jane Doe takes metformin."), "pdf")
    pdf = PdfReader(io.BytesIO(base64.b64decode(result["data"])))
    assert "[NAME] takes metformin." in "".join(
        page.extract_text() for page in pdf.pages
    )
    assert "Jane Doe" not in str(pdf.metadata)
    assert result["entity_count"] == 1
    assert result["entities"][0]["location"] == {"page": 1}
    assert "layout" in result["warnings"][0]


def test_pdf_without_optional_extra_is_a_clear_validation_error(monkeypatch):
    pdf = make_pdf()
    monkeypatch.setitem(__import__("sys").modules, "pypdf", None)
    with pytest.raises(ValueError, match=r"outis\[pdf\]"):
        invoke(pdf, "pdf")


def test_blank_encrypted_attached_and_annotated_pdfs_rejected():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (
        ArrayObject,
        DictionaryObject,
        NameObject,
        TextStringObject,
    )

    with pytest.raises(ValueError, match="OCR"):
        invoke(make_pdf(blank=True), "pdf")
    for feature in ["encrypted", "attached", "annotation"]:
        writer = PdfWriter()
        writer.append_pages_from_reader(PdfReader(io.BytesIO(make_pdf())))
        if feature == "encrypted":
            writer.encrypt("password")
        elif feature == "attached":
            writer.add_attachment("patient.txt", b"Jane Doe")
        else:
            writer.pages[0][NameObject("/Annots")] = ArrayObject(
                [
                    DictionaryObject(
                        {NameObject("/Contents"): TextStringObject("Jane Doe")}
                    )
                ]
            )
        buf = io.BytesIO()
        writer.write(buf)
        with pytest.raises(ValueError):
            invoke(buf.getvalue(), "pdf")


@pytest.mark.parametrize(
    "kind,part,xml",
    [
        (
            "docx",
            "word/styles.xml",
            f'<w:styles xmlns:w="{W}"><w:style w:styleId="JaneDoe"><w:name w:val="Jane Doe"/></w:style></w:styles>',
        ),
        (
            "xlsx",
            "xl/styles.xml",
            f'<styleSheet xmlns="{S}"><numFmts><numFmt numFmtId="166" formatCode="&quot;Jane Doe&quot; 0"/></numFmts></styleSheet>',
        ),
        (
            "xlsx",
            "xl/styles.xml",
            f'<styleSheet xmlns="{S}"><cellStyles><cellStyle name="Jane Doe"/></cellStyles></styleSheet>',
        ),
    ],
)
def test_identifiers_in_unsupported_office_attributes_fail_closed(kind, part, xml):
    with pytest.raises(ValueError, match="unsupported attributes"):
        invoke(make_office(kind, {part: xml}), kind)


def test_style_identifiers_and_bookmarks_are_consistently_remapped():
    parts = {
        "word/styles.xml": f'<w:styles xmlns:w="{W}"><w:style w:styleId="JaneDoe"><w:name w:val="Custom"/></w:style></w:styles>',
        "word/document.xml": f'<w:document xmlns:w="{W}"><w:body><w:p><w:pPr><w:pStyle w:val="JaneDoe"/></w:pPr><w:bookmarkStart w:name="JaneDoe" w:id="0"/><w:r><w:t>Jane Doe</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p></w:body></w:document>',
    }
    members = unpack(invoke(make_office("docx", parts), "docx"))
    assert all(
        b"Jane" not in value and b"Doe" not in value for value in members.values()
    )
    document = ET.fromstring(members["word/document.xml"])
    assert (
        next(document.iter(f"{{{W}}}bookmarkStart")).get(f"{{{W}}}name") == "bookmark_0"
    )
    style = ET.fromstring(members["word/styles.xml"])
    assert next(document.iter(f"{{{W}}}pStyle")).get(f"{{{W}}}val") == next(
        style.iter(f"{{{W}}}style")
    ).get(f"{{{W}}}styleId")


def test_unclaimed_text_is_transformed_and_auxiliary_parts_are_removed():
    parts = {
        "customXml/item1.xml": "<bibliography>Jane Doe</bibliography>",
        "word/settings.xml": f'<w:settings xmlns:w="{W}"><w:extension>Jane Doe</w:extension></w:settings>',
    }
    result = invoke(make_office("docx", parts), "docx")
    members = unpack(result)
    assert "customXml/item1.xml" not in members
    assert b"[NAME]" in members["word/settings.xml"]
    assert all(b"Jane" not in value for value in members.values())


def test_deep_office_xml_rejected():
    with pytest.raises(ValueError, match="unsafe XML"):
        invoke(
            make_office(
                "docx", {"word/settings.xml": "<x>" * 65 + "Jane Doe" + "</x>" * 65}
            ),
            "docx",
        )


def test_pdf_compressed_content_is_bounded_before_extraction():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    from outis.formats.files import MAX_ZIP_BYTES

    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(io.BytesIO(make_pdf())))
    stream = DecodedStreamObject()
    stream.set_data(b" " * (MAX_ZIP_BYTES + 1))
    writer.pages[0][NameObject("/Contents")] = writer._add_object(stream.flate_encode())
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ValueError, match="20 MiB"):
        invoke(buf.getvalue(), "pdf")


def test_pdf_unicode_english_text_remains_extractable():
    from pathlib import Path

    import reportlab
    from pypdf import PdfReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    pdfmetrics.registerFont(
        TTFont("TestVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"))
    )
    buf = io.BytesIO()
    canvas = Canvas(buf)
    canvas.setFont("TestVera", 12)
    canvas.drawString(20, 750, "Jane Doe visited the café.")
    canvas.save()
    result = invoke(buf.getvalue(), "pdf")
    reader = PdfReader(io.BytesIO(base64.b64decode(result["data"])))
    assert "[NAME] visited the café." in reader.pages[0].extract_text()


def test_encoded_word_symbol_text_rejected():
    symbols = "".join(
        f'<w:sym w:font="Arial" w:char="{ord(char):04X}"/>' for char in "Jane Doe"
    )
    document = f'<w:document xmlns:w="{W}"><w:body><w:p><w:r>{symbols}</w:r></w:p></w:body></w:document>'
    with pytest.raises(ValueError, match="unsupported fields"):
        invoke(make_office("docx", {"word/document.xml": document}), "docx")


def test_word_hyperlinks_rejected_to_avoid_losing_clickable_labels():
    document = f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Jane Doe</w:t></w:r><w:hyperlink w:anchor="JaneDoe"><w:r><w:t>Jump</w:t></w:r></w:hyperlink></w:p></w:body></w:document>'
    with pytest.raises(ValueError, match="unsupported fields"):
        invoke(make_office("docx", {"word/document.xml": document}), "docx")


def test_escaped_spreadsheet_display_literals_rejected():
    styles = f'<styleSheet xmlns="{S}"><numFmts><numFmt numFmtId="166" formatCode="\\J\\a\\n\\e \\D\\o\\e 0"/></numFmts></styleSheet>'
    with pytest.raises(ValueError, match="literal or escaped"):
        invoke(make_office("xlsx", {"xl/styles.xml": styles}), "xlsx")


@pytest.mark.parametrize(
    "payload", ["_x004A_ane Doe", "Jane _x0044_oe", "_x005F_x004A_ane Doe"]
)
def test_office_encoded_escape_sequences_fail_closed(payload):
    sheet = f'<worksheet xmlns="{S}"><sheetData><row><c r="A1" t="inlineStr"><is><t>{payload}</t></is></c></row></sheetData></worksheet>'
    with pytest.raises(ValueError, match="escaped character"):
        invoke(make_office("xlsx", {"xl/worksheets/sheet1.xml": sheet}), "xlsx")


def test_presentation_custom_path_geometry_rejected():
    slide = f'<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:spPr><a:custGeom><a:pathLst/></a:custGeom></p:spPr></p:sp></p:spTree></p:cSld></p:sld>'
    with pytest.raises(ValueError, match="unsupported fields"):
        invoke(make_office("pptx", {"ppt/slides/slide1.xml": slide}), "pptx")


RELATIONSHIPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def test_package_relationship_ids_targets_and_types_are_rebuilt():
    parts = {
        "_rels/.rels": f'<Relationships xmlns="{R}"><Relationship Id="Jane_Doe" Type="{RELATIONSHIPS}/officeDocument" Target="word/JaneDoe/../document.xml" TargetMode="Internal">JaneDoe</Relationship></Relationships>',
        "word/_rels/document.xml.rels": f'<Relationships xmlns="{R}"><Relationship Id="JaneDoeStyles" Type="{RELATIONSHIPS}/styles" Target="styles.xml"/><Relationship Id="JaneDoeHeader" Type="{RELATIONSHIPS}/header" Target="header1.xml"/></Relationships>',
        "word/styles.xml": f'<w:styles xmlns:w="{W}"/>',
        "word/header1.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Header</w:t></w:r></w:p></w:hdr>',
        "word/document.xml": f'<w:document xmlns:w="{W}" xmlns:r="{RELATIONSHIPS}"><w:body><w:p><w:r><w:t>Seen</w:t></w:r></w:p><w:sectPr><w:headerReference w:type="default" r:id="JaneDoeHeader"/></w:sectPr></w:body></w:document>',
        "[Content_Types].xml": f'<Types xmlns="{C}"><Default Extension="JaneDoe" ContentType="application/xml"/><Default Extension="XML" ContentType="application/xml"/><Override PartName="/WORD/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/JaneDoe.xml" ContentType="application/xml"/></Types>',
    }
    members = unpack(invoke(make_office("docx", parts), "docx"))
    assert all(
        b"Jane" not in value and b"Doe" not in value for value in members.values()
    )
    package = ET.fromstring(members["_rels/.rels"])
    assert [rel.attrib for rel in package] == [
        {
            "Id": "rId1",
            "Type": f"{RELATIONSHIPS}/officeDocument",
            "Target": "word/document.xml",
        }
    ]
    links = {
        rel.get("Target"): rel.get("Id")
        for rel in ET.fromstring(members["word/_rels/document.xml.rels"])
    }
    assert links == {"styles.xml": "rId1", "header1.xml": "rId2"}
    document = ET.fromstring(members["word/document.xml"])
    header = next(document.iter(f"{{{W}}}headerReference"))
    assert header.get(f"{{{RELATIONSHIPS}}}id") == links["header1.xml"]
    types = ET.fromstring(members["[Content_Types].xml"])
    assert [item.attrib for item in types] == [
        {"Extension": "xml", "ContentType": "application/xml"},
        {
            "PartName": "/word/document.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        },
    ]


def test_presentation_relationship_references_stay_consistent():
    parts = {
        "ppt/_rels/presentation.xml.rels": f'<Relationships xmlns="{R}"><Relationship Id="JaneDoe" Type="{RELATIONSHIPS}/slide" Target="slides/slide1.xml"/></Relationships>',
        "ppt/presentation.xml": f'<p:presentation xmlns:p="{P}" xmlns:r="{RELATIONSHIPS}"><p:sldIdLst><p:sldId id="256" r:id="JaneDoe"/></p:sldIdLst></p:presentation>',
    }
    members = unpack(invoke(make_office("pptx", parts), "pptx"))
    assert all(b"Jane" not in value for value in members.values())
    slide_id = next(
        ET.fromstring(members["ppt/presentation.xml"]).iter(f"{{{P}}}sldId")
    )
    rel = next(iter(ET.fromstring(members["ppt/_rels/presentation.xml.rels"])))
    assert slide_id.get(f"{{{RELATIONSHIPS}}}id") == rel.get("Id") == "rId1"
    assert rel.get("Target") == "slides/slide1.xml"


@pytest.mark.parametrize(
    "part,xml",
    [
        (
            "_rels/.rels",
            f'<Relationships xmlns="{R}"><Relationship Id="r1" Type="urn:Jane Doe" Target="word/document.xml"/></Relationships>',
        ),
        (
            "[Content_Types].xml",
            f'<Types xmlns="{C}"><Override PartName="/word/document.xml" ContentType="application/Jane Doe+xml"/></Types>',
        ),
    ],
)
def test_identifiers_in_unrecognized_package_types_fail_closed(part, xml):
    with pytest.raises(ValueError, match="unsupported attributes"):
        invoke(make_office("docx", {part: xml}), "docx")


def test_unresolved_relationship_reference_rejected():
    part = f'<w:document xmlns:w="{W}" xmlns:r="{RELATIONSHIPS}"><w:body><w:sectPr><w:headerReference r:id="JaneDoe"/></w:sectPr></w:body></w:document>'
    with pytest.raises(ValueError, match="unresolved relationship"):
        invoke(make_office("docx", {"word/document.xml": part}), "docx")


# The fixtures above are deliberately skeletal parser inputs. The tests below
# start from complete documents written by python-docx, openpyxl and
# python-pptx, and reopen the returned bytes with the same independent readers.


def assert_complete_package(raw):
    """Every internal relationship and content-type override names a real part."""
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = set(archive.namelist())
        missing = []
        for name in names:
            if not name.endswith(".rels"):
                continue
            source_dir = posixpath.dirname(posixpath.dirname(name))
            for relationship in ET.fromstring(archive.read(name)):
                if relationship.get("TargetMode") == "External":
                    continue
                target = relationship.get("Target")
                part = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join(source_dir, target))
                )
                if part not in names:
                    missing.append((name, target))
        for override in ET.fromstring(archive.read("[Content_Types].xml")):
            part = override.get("PartName")
            if part and part.lstrip("/") not in names:
                missing.append(("[Content_Types].xml", part))
    assert missing == []


def saved(document):
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def redacted(raw, kind):
    result = invoke(raw, kind)
    output = base64.b64decode(result["data"])
    assert_complete_package(output)
    with zipfile.ZipFile(io.BytesIO(output)) as archive:
        assert all(
            b"Jane" not in archive.read(name) and b"Doe" not in archive.read(name)
            for name in archive.namelist()
        )
    return result, io.BytesIO(output)


def test_valid_docx_round_trips_through_an_independent_reader():
    docx = pytest.importorskip("docx")

    document = docx.Document()
    document.core_properties.author = "Jane Doe"
    paragraph = document.add_paragraph("Seen today: ")
    paragraph.add_run("Jane ")
    paragraph.add_run("Doe").bold = True
    paragraph.add_run(" takes metformin.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Patient"
    table.cell(0, 1).text = "Jane Doe"
    document.sections[0].header.paragraphs[0].text = "Chart for Jane Doe"
    document.sections[0].footer.paragraphs[0].text = "Page footer"
    result, output = redacted(saved(document), "docx")

    reopened = docx.Document(output)
    assert [p.text for p in reopened.paragraphs] == [
        "Seen today: [NAME] takes metformin."
    ]
    table = reopened.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Patient", "[NAME]"]
    section = reopened.sections[0]
    assert section.header.paragraphs[0].text == "Chart for [NAME]"
    assert section.footer.paragraphs[0].text == "Page footer"
    assert reopened.core_properties.author == ""
    assert result["entity_count"] == 3


def test_valid_xlsx_round_trips_through_an_independent_reader():
    from datetime import date

    openpyxl = pytest.importorskip("openpyxl")

    workbook = openpyxl.Workbook()
    workbook.properties.creator = "Jane Doe"
    sheet = workbook.active
    sheet.title = "Jane Doe"
    sheet.append(["Patient", "Visit", "Glucose"])
    sheet.append(["Jane Doe", date(2023, 3, 15), 109])
    second = workbook.create_sheet("Labs")
    second["A1"] = "Reviewed with Jane Doe"
    result, output = redacted(saved(workbook), "xlsx")

    reopened = openpyxl.load_workbook(output)
    assert len(reopened.sheetnames) == 2
    assert all("Jane" not in name for name in reopened.sheetnames)
    first, labs = (reopened[name] for name in reopened.sheetnames)
    assert [[cell.value for cell in row] for row in first.iter_rows()] == [
        ["Patient", "Visit", "Glucose"],
        ["[NAME]", "2023-03-15", "109"],
    ]
    assert labs["A1"].value == "Reviewed with [NAME]"
    assert reopened.properties.creator != "Jane Doe"
    assert result["entity_count"] == 3


def test_valid_pptx_round_trips_through_an_independent_reader():
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    presentation.core_properties.author = "Jane Doe"
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Jane Doe follow-up"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    for text in ["Seen by ", "Jane ", "Doe", " today."]:
        paragraph.add_run().text = text
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Plan"
    result, output = redacted(saved(presentation), "pptx")

    reopened = pptx.Presentation(output)
    assert [
        [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        for slide in reopened.slides
    ] == [["[NAME] follow-up", "Seen by [NAME] today."], ["Plan"]]
    assert reopened.core_properties.author == ""
    assert result["entity_count"] == 2
